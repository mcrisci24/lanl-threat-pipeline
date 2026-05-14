from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor


PROJECT_ROOT = Path(__file__).resolve().parent

INPUT_PPTX = PROJECT_ROOT / "presentation_ml_final" / "LANL_ML_Final_Project_Presentation.pptx"
OUTPUT_PPTX = PROJECT_ROOT / "presentation_ml_final" / "LANL_ML_Final_Project_Presentation_with_diagnostics.pptx"

ASSET_DIR_CANDIDATES = [
    PROJECT_ROOT / "presentation_ml_final" / "assets" / "lgbm_diagnostics",
    PROJECT_ROOT / "reports" / "figures" / "lgbm_diagnostics",
]

# Insert after current Slide 9, which is the model results slide.
# New slides become Slide 10 and Slide 11.
INSERT_AFTER_SLIDE_NUMBER = 9


NAVY = RGBColor(11, 31, 59)
TEAL = RGBColor(0, 150, 136)
DARK = RGBColor(30, 30, 30)
GRAY = RGBColor(90, 90, 90)
LIGHT_GRAY = RGBColor(245, 247, 250)
WHITE = RGBColor(255, 255, 255)


def find_asset_dir():
    for folder in ASSET_DIR_CANDIDATES:
        if folder.exists():
            required = [
                "lgbm_roc_validation.png",
                "lgbm_roc_test.png",
                "lgbm_pr_validation.png",
                "lgbm_pr_test.png",
                "lgbm_confusion_matrix_validation.png",
                "lgbm_confusion_matrix_test.png",
            ]
            missing = [name for name in required if not (folder / name).exists()]
            if not missing:
                return folder

    checked = "\n".join(str(p) for p in ASSET_DIR_CANDIDATES)
    raise FileNotFoundError(
        "Could not find all LightGBM diagnostic images.\n"
        "Checked folders:\n"
        f"{checked}\n\n"
        "Run:\n"
        "py -3 make_lgbm_diagnostics.py --threshold 0.10 "
        "--predictions model_outputs/lightgbm_model/validation_test_predictions.csv"
    )


def add_textbox(slide, text, left, top, width, height, font_size=18, bold=False, color=DARK):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = text
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_title(slide, title, subtitle=None):
    add_textbox(slide, title, 0.45, 0.20, 12.4, 0.45, font_size=28, bold=True, color=NAVY)

    if subtitle:
        add_textbox(slide, subtitle, 0.48, 0.72, 12.2, 0.32, font_size=13, bold=False, color=GRAY)


def add_footer(slide, slide_number):
    add_textbox(
        slide,
        f"Mark Crisci  |  LANL ML Final Project  |  {slide_number}/14",
        0.50,
        7.05,
        6.5,
        0.25,
        font_size=9,
        color=GRAY,
    )


def add_bullet_box(slide, bullets, left, top, width, height, title=None):
    shape = slide.shapes.add_shape(
        1,
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(height),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = LIGHT_GRAY
    shape.line.color.rgb = RGBColor(220, 225, 230)

    text_box = slide.shapes.add_textbox(
        Inches(left + 0.15),
        Inches(top + 0.10),
        Inches(width - 0.30),
        Inches(height - 0.20),
    )
    tf = text_box.text_frame
    tf.clear()

    if title:
        p = tf.paragraphs[0]
        run = p.add_run()
        run.text = title
        run.font.size = Pt(14)
        run.font.bold = True
        run.font.color.rgb = NAVY
        start_index = 1
    else:
        start_index = 0

    for i, bullet in enumerate(bullets):
        if title or i > 0:
            p = tf.add_paragraph()
        else:
            p = tf.paragraphs[0]

        p.text = bullet
        p.level = 0
        p.font.size = Pt(10.5)
        p.font.color.rgb = DARK

    return text_box


def blank_slide(prs):
    layout = prs.slide_layouts[6]
    return prs.slides.add_slide(layout)


def add_curves_slide(prs, asset_dir):
    slide = blank_slide(prs)

    add_title(
        slide,
        "Winning Model Diagnostics: ROC and PR Curves",
        "Validation and test curves for the LightGBM rare-event classifier",
    )

    # 2x2 grid
    slide.shapes.add_picture(str(asset_dir / "lgbm_roc_validation.png"), Inches(0.45), Inches(1.10), width=Inches(3.15))
    slide.shapes.add_picture(str(asset_dir / "lgbm_roc_test.png"), Inches(3.75), Inches(1.10), width=Inches(3.15))
    slide.shapes.add_picture(str(asset_dir / "lgbm_pr_validation.png"), Inches(0.45), Inches(4.02), width=Inches(3.15))
    slide.shapes.add_picture(str(asset_dir / "lgbm_pr_test.png"), Inches(3.75), Inches(4.02), width=Inches(3.15))

    bullets = [
        "ROC-AUC shows global ranking ability: LightGBM ranks attack-like host-hours above benign ones.",
        "PR-AUC is the primary metric because only 0.0043% of rows are positives.",
        "Test ROC-AUC = 0.881 and test PR-AUC = 0.03905.",
        "The PR curve looks visually compressed because the random baseline is only 0.000043.",
    ]
    add_bullet_box(
        slide,
        bullets,
        left=7.15,
        top=1.20,
        width=5.75,
        height=4.75,
        title="How to interpret this slide",
    )

    add_textbox(
        slide,
        "Main takeaway: the model is not a perfect alarm. It is a risk-ranking system that moves rare attack behavior far above random.",
        7.25,
        6.15,
        5.55,
        0.55,
        font_size=12,
        bold=True,
        color=TEAL,
    )

    add_footer(slide, 10)
    return slide


def add_confusion_slide(prs, asset_dir):
    slide = blank_slide(prs)

    add_title(
        slide,
        "Winning Model Diagnostics: Confusion Matrices",
        "Validation and test operating-point behavior at threshold = 0.10",
    )

    slide.shapes.add_picture(
        str(asset_dir / "lgbm_confusion_matrix_validation.png"),
        Inches(0.50),
        Inches(1.10),
        width=Inches(5.15),
    )

    slide.shapes.add_picture(
        str(asset_dir / "lgbm_confusion_matrix_test.png"),
        Inches(5.95),
        Inches(1.10),
        width=Inches(5.15),
    )

    bullets = [
        "Validation: TP = 23, FN = 96, FP = 57,773.",
        "Test: TP = 28, FN = 91, FP = 57,535.",
        "At threshold 0.10, the test model catches 28 of 119 attacks, about 23.5% recall.",
        "The false positives are the analyst workload. Raising the threshold reduces noise but misses more attacks.",
    ]

    add_bullet_box(
        slide,
        bullets,
        left=11.15,
        top=1.15,
        width=1.95,
        height=5.35,
        title="Operational reading",
    )

    add_textbox(
        slide,
        "This is why the threshold is a business decision: lower threshold = more recall, more alerts; higher threshold = fewer alerts, more missed attacks.",
        0.60,
        6.70,
        12.1,
        0.35,
        font_size=11,
        bold=True,
        color=TEAL,
    )

    add_footer(slide, 11)
    return slide


def move_slide(prs, old_index, new_index):
    """
    Move a slide inside the presentation using python-pptx internals.
    Indexes are zero-based.
    """
    slides = prs.slides._sldIdLst
    slide_ids = list(slides)
    slide_to_move = slide_ids[old_index]
    slides.remove(slide_to_move)
    slides.insert(new_index, slide_to_move)


def main():
    if not INPUT_PPTX.exists():
        raise FileNotFoundError(f"Missing input PowerPoint: {INPUT_PPTX}")

    asset_dir = find_asset_dir()
    print(f"[INFO] Using diagnostic images from: {asset_dir}")

    prs = Presentation(str(INPUT_PPTX))

    original_count = len(prs.slides)
    print(f"[INFO] Original slide count: {original_count}")

    # Add curves slide at end, then move it after slide 9.
    add_curves_slide(prs, asset_dir)
    move_slide(prs, len(prs.slides) - 1, INSERT_AFTER_SLIDE_NUMBER)

    # Add confusion slide at end, then move it after the curves slide.
    add_confusion_slide(prs, asset_dir)
    move_slide(prs, len(prs.slides) - 1, INSERT_AFTER_SLIDE_NUMBER + 1)

    prs.save(str(OUTPUT_PPTX))

    print(f"[INFO] Saved: {OUTPUT_PPTX}")
    print(f"[INFO] New slide count: {len(prs.slides)}")
    print("[INFO] Open the new deck and check slides 10 and 11.")


if __name__ == "__main__":
    main()